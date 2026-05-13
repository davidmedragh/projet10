"""
Script de génération du support de présentation — Étape 2
Projet 10 — Labellisez et appliquez des approches semi-supervisées en traitement d'images

Auteur : David MEDRAGH
Date : Mai 2026

Description :
    Ce script génère un fichier PowerPoint (.pptx) résumant les résultats de
    l'extraction de features réalisée dans le notebook
    projet10_etape2_features.ipynb.

    Le template visuel reprend le style du projet 9 (palette bleu profond / or,
    barre de titre numérotée, footer avec logos, KPI cards, séparateurs de section).

Slides produites :
    1. Couverture
    2. Sommaire
    3. Séparateur PRÉTRAITEMENT
    4. Pipeline de preprocessing
    5. Vérification visuelle (brute vs preprocessed)
    6. Séparateur EXTRACTION
    7. Modèle ResNet50 — Configuration
    8. Résultats de l'extraction (KPI cards)
    9. Tableau exploitable
   10. Séparateur BILAN
   11. Synthèse & Prochaines étapes

Usage :
    uv run python scripts/generate_pptx_etape2.py

Sortie :
    livrables/projet10_etape2_features.pptx
"""

import warnings
warnings.filterwarnings("ignore")

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- CHEMINS ---
PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT_ROOT / "livrables"
OUTPUT_FILE = OUTPUT_DIR / "projet10_etape2_features.pptx"
LOGO_CURELYTICSAI = PROJECT_ROOT / "doc" / "logo" / "logo_CurelyticsIA.png"
LOGO_OC = PROJECT_ROOT / "doc" / "logo" / "logo_openclassrooms.png"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# --- PALETTE ---
BLEU_PROFOND = RGBColor(0x0D, 0x1B, 0x2A)
BLEU_PRINCIPAL = RGBColor(0x1B, 0x3A, 0x5C)
BLEU_ACCENT = RGBColor(0x3A, 0x7C, 0xBD)
OR_ACCENT = RGBColor(0xD4, 0xA0, 0x3C)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
GRIS_98 = RGBColor(0xFA, 0xFA, 0xFA)
GRIS_TEXTE = RGBColor(0x2D, 0x2D, 0x2D)
GRIS_SUBTITLE = RGBColor(0x6B, 0x6B, 0x6B)
GRIS_FOOTER = RGBColor(0xAA, 0xAA, 0xAA)
GRIS_LIGNE = RGBColor(0xDD, 0xDD, 0xDD)
VERT_OK = RGBColor(0x27, 0xAE, 0x60)
ORANGE_WARN = RGBColor(0xE6, 0x7E, 0x22)


# ──────────────────────────────────────────────
# Helpers (même pattern que projet 9 / étape 1)
# ──────────────────────────────────────────────

def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.fill.solid()
        shape.line.fill.fore_color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf


def set_para(tf, text, size=Pt(14), color=GRIS_TEXTE, bold=False,
             align=PP_ALIGN.LEFT, space_before=Pt(0), space_after=Pt(0), new=False):
    if new:
        p = tf.add_paragraph()
    else:
        p = (tf.paragraphs[0]
             if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == ""
             else tf.add_paragraph())
    run = p.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    p.alignment = align
    p.space_before = space_before
    p.space_after = space_after
    return p


def slide_title_bar(slide, number, title):
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.3), BLEU_PROFOND)
    add_rect(slide, 0, 0, Inches(0.08), Inches(1.3), OR_ACCENT)
    tf = add_textbox(slide, Inches(0.5), Inches(0.25), Inches(0.8), Inches(0.8))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf, f"{number:02d}", size=Pt(36), color=OR_ACCENT, bold=True)
    add_rect(slide, Inches(1.35), Inches(0.3), Pt(2), Inches(0.7), OR_ACCENT)
    tf = add_textbox(slide, Inches(1.6), Inches(0.25), Inches(10), Inches(0.8))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf, title, size=Pt(26), color=BLANC, bold=True)


def slide_footer(slide):
    add_rect(slide, 0, Inches(7.1), SLIDE_W, Pt(1), GRIS_LIGNE)
    tf = add_textbox(slide, Inches(0.5), Inches(7.15), Inches(6), Inches(0.3))
    set_para(tf, "CurelyticsIA — Extraction de features IRM", size=Pt(9), color=GRIS_FOOTER)
    tf2 = add_textbox(slide, Inches(6.5), Inches(7.15), Inches(2), Inches(0.3))
    set_para(tf2, "Projet 10", size=Pt(9), color=BLEU_ACCENT, bold=True,
             align=PP_ALIGN.CENTER)
    tf3 = add_textbox(slide, Inches(9), Inches(7.15), Inches(3.8), Inches(0.3))
    set_para(tf3, "David Medragh — Data Scientist", size=Pt(9), color=GRIS_FOOTER,
             align=PP_ALIGN.RIGHT)
    logo_h = Inches(0.3)
    if LOGO_OC.exists():
        slide.shapes.add_picture(str(LOGO_OC), Inches(4.8), Inches(7.13), height=logo_h)
    if LOGO_CURELYTICSAI.exists():
        slide.shapes.add_picture(str(LOGO_CURELYTICSAI), Inches(5.5), Inches(7.13), height=logo_h)


def add_content_block(slide, title, bullets, left=Inches(0.8), top=Inches(1.6),
                      width=Inches(5.5)):
    tf = add_textbox(slide, left, top, width, Inches(5))
    set_para(tf, title, size=Pt(16), color=BLEU_PRINCIPAL, bold=True, space_after=Pt(8))
    for bullet in bullets:
        set_para(tf, f"▸  {bullet}", size=Pt(13), color=GRIS_TEXTE, space_before=Pt(5), new=True)
    return tf


def add_kpi_card(slide, left, top, value, label, value_color=BLEU_ACCENT,
                 w=Inches(2.5), h=Inches(1.3)):
    add_rect(slide, left, top, w, h, GRIS_98, GRIS_LIGNE)
    tf = add_textbox(slide, left + Inches(0.15), top + Inches(0.1),
                     w - Inches(0.3), Inches(0.7))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf, value, size=Pt(28), color=value_color, bold=True, align=PP_ALIGN.CENTER)
    tf2 = add_textbox(slide, left + Inches(0.15), top + Inches(0.75),
                      w - Inches(0.3), Inches(0.5))
    set_para(tf2, label, size=Pt(11), color=GRIS_SUBTITLE, align=PP_ALIGN.CENTER)


def add_check_item(tf, text, ok=True):
    prefix = "✓" if ok else "⚠"
    color = VERT_OK if ok else ORANGE_WARN
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = f"  {prefix}  {text}"
    run.font.size = Pt(14)
    run.font.color.rgb = color
    p.space_before = Pt(6)


def add_table(slide, headers, rows, left=Inches(0.8), top=Inches(1.8), width=Inches(11.5)):
    nb_rows = len(rows) + 1
    nb_cols = len(headers)
    table_shape = slide.shapes.add_table(
        nb_rows, nb_cols, left, top, width, Inches(0.38 * nb_rows))
    table = table_shape.table

    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLEU_PROFOND
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(12)
            paragraph.font.bold = True
            paragraph.font.color.rgb = BLANC
            paragraph.font.name = "Calibri"
            paragraph.alignment = PP_ALIGN.CENTER

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLANC if i % 2 == 0 else GRIS_98
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(11)
                paragraph.font.color.rgb = GRIS_TEXTE
                paragraph.font.name = "Calibri"
                paragraph.alignment = PP_ALIGN.CENTER


# ──────────────────────────────────────────────
# Construction des slides
# ──────────────────────────────────────────────

def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ═══════ SLIDE 1 — COUVERTURE ═══════
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, BLEU_PROFOND)
    add_rect(s, 0, 0, SLIDE_W, Inches(0.06), OR_ACCENT)
    add_rect(s, 0, Inches(7.44), SLIDE_W, Inches(0.06), OR_ACCENT)

    tf = add_textbox(s, Inches(1.2), Inches(1.8), Inches(10), Inches(1.5))
    set_para(tf, "EXTRACTION DE FEATURES", size=Pt(44), color=BLANC, bold=True)
    set_para(tf, "IRM CÉRÉBRALES", size=Pt(44), color=OR_ACCENT, bold=True, new=True)

    add_rect(s, Inches(1.2), Inches(3.8), Inches(3), Pt(3), OR_ACCENT)

    tf2 = add_textbox(s, Inches(1.2), Inches(4.1), Inches(10), Inches(1))
    set_para(tf2, "Étape 2 — Prétraitez et extrayez les features",
             size=Pt(20), color=RGBColor(0xBB, 0xBB, 0xBB))
    set_para(tf2, "Projet 10 — Labellisation et approches semi-supervisées",
             size=Pt(15), color=RGBColor(0x88, 0x88, 0x88), new=True, space_before=Pt(6))

    tf3 = add_textbox(s, Inches(1.2), Inches(5.5), Inches(10), Inches(1))
    set_para(tf3, "David MEDRAGH", size=Pt(18), color=BLANC, bold=True)
    set_para(tf3, "Data Scientist — CurelyticsIA", size=Pt(14),
             color=RGBColor(0x88, 0x88, 0x88), new=True, space_before=Pt(4))
    set_para(tf3, "Mai 2026", size=Pt(12),
             color=RGBColor(0x66, 0x66, 0x66), new=True, space_before=Pt(12))

    logo_h = Inches(0.6)
    if LOGO_CURELYTICSAI.exists():
        s.shapes.add_picture(str(LOGO_CURELYTICSAI), Inches(10.5), Inches(6.2), height=logo_h)
    if LOGO_OC.exists():
        s.shapes.add_picture(str(LOGO_OC), Inches(11.8), Inches(6.2), height=logo_h)

    # ═══════ SLIDE 2 — SOMMAIRE ═══════
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, Inches(1.3), BLEU_PROFOND)
    add_rect(s, 0, 0, Inches(0.08), Inches(1.3), OR_ACCENT)
    tf = add_textbox(s, Inches(0.5), Inches(0.25), Inches(10), Inches(0.8))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf, "Sommaire", size=Pt(26), color=BLANC, bold=True)

    x_positions = [Inches(0.5), Inches(4.8), Inches(9.2)]
    sommaire_sections = [
        ("Prétraitement", "Images → tenseurs", [
            "Pipeline de preprocessing officiel",
            "Vérification visuelle",
        ]),
        ("Extraction", "ResNet50 → features", [
            "Configuration du modèle",
            "Résultats de l'extraction",
            "Tableau exploitable",
        ]),
        ("Bilan", "Résultats et suite", [
            "Synthèse & Prochaines étapes",
        ]),
    ]

    for col, (partie, subtitle, items) in enumerate(sommaire_sections):
        x = x_positions[col]
        y_start = Inches(1.6)
        tf = add_textbox(s, x, y_start, Inches(4), Inches(0.5))
        set_para(tf, partie, size=Pt(18), color=OR_ACCENT, bold=True)
        set_para(tf, subtitle, size=Pt(12), color=GRIS_SUBTITLE, new=True,
                 space_before=Pt(4))
        tf2 = add_textbox(s, x + Inches(0.2), y_start + Inches(0.9), Inches(3.8),
                          Inches(4.5))
        for i, item in enumerate(items):
            set_para(tf2, f"▸  {item}", size=Pt(11), color=GRIS_TEXTE,
                     space_before=Pt(6), new=(i > 0))

    slide_footer(s)

    # ═══════ SLIDE 3 — SÉPARATEUR PRÉTRAITEMENT ═══════
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, BLEU_PROFOND)
    add_rect(s, 0, 0, SLIDE_W, Inches(0.06), OR_ACCENT)
    tf = add_textbox(s, Inches(1.2), Inches(2.5), Inches(10), Inches(1))
    set_para(tf, "PRÉTRAITEMENT", size=Pt(44), color=OR_ACCENT, bold=True)
    tf2 = add_textbox(s, Inches(1.2), Inches(3.8), Inches(10), Inches(1))
    set_para(tf2, "Des images brutes aux tenseurs normalisés", size=Pt(24), color=BLANC)
    set_para(tf2, "Pipeline officiel ResNet50 — Resize, CenterCrop, Normalize",
             size=Pt(14), color=RGBColor(0x88, 0x88, 0x88), new=True, space_before=Pt(12))

    # ═══════ SLIDE 4 — PIPELINE DE PREPROCESSING ═══════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 1, "Pipeline de preprocessing")

    add_content_block(s, "Preprocessing officiel ResNet50", [
        "weights = ResNet50_Weights.DEFAULT",
        "preprocess = weights.transforms()",
        "Pipeline unique — pas de transforms manuel parallèle",
    ], left=Inches(0.8), top=Inches(1.6), width=Inches(5.5))

    add_content_block(s, "Étapes du pipeline", [
        "Resize(232) — redimensionnement du côté court",
        "CenterCrop(224) — recadrage central 224×224",
        "ToTensor() — conversion en tenseur [0, 1]",
        "Normalize(mean=[.485, .456, .406], std=[.229, .224, .225])",
    ], left=Inches(7.0), top=Inches(1.6), width=Inches(5.5))

    add_content_block(s, "Dataset unifié", [
        "1 506 images chargées dans un DataFrame unique",
        "6 colonnes de métadonnées : path, filename, dossier, label_name, label_id, is_labeled",
        "DataLoader avec shuffle=False — ordre stable garanti",
        "Batch size = 32 → 48 batches",
    ], left=Inches(0.8), top=Inches(4.0), width=Inches(11))

    slide_footer(s)

    # ═══════ SLIDE 5 — VÉRIFICATION VISUELLE ═══════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 2, "Vérification visuelle — Brute vs Preprocessed")

    add_content_block(s, "Méthode", [
        "Colonne gauche : image brute chargée avec OpenCV (cv2.imread + cvtColor BGR→RGB)",
        "Colonne droite : image après weights.transforms() puis dénormalisation ImageNet",
        "3 catégories vérifiées : cancer, normal, sans_label",
    ], left=Inches(0.8), top=Inches(1.6), width=Inches(11))

    add_content_block(s, "Outils de l'école utilisés", [
        "OpenCV (cv2) — recommandé par l'école, utilisé pour le chargement brut et la conversion BGR→RGB",
        "torchvision — preprocessing officiel du modèle via weights.transforms()",
        "PIL (Pillow) — ouverture des images dans le Dataset custom",
        "matplotlib — affichage comparatif",
    ], left=Inches(0.8), top=Inches(3.8), width=Inches(11))

    slide_footer(s)

    # ═══════ SLIDE 6 — SÉPARATEUR EXTRACTION ═══════
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, BLEU_PROFOND)
    add_rect(s, 0, 0, SLIDE_W, Inches(0.06), OR_ACCENT)
    tf = add_textbox(s, Inches(1.2), Inches(2.5), Inches(10), Inches(1))
    set_para(tf, "EXTRACTION", size=Pt(44), color=OR_ACCENT, bold=True)
    tf2 = add_textbox(s, Inches(1.2), Inches(3.8), Inches(10), Inches(1))
    set_para(tf2, "Des tenseurs aux embeddings visuels", size=Pt(24), color=BLANC)
    set_para(tf2, "ResNet50 pré-entraîné — 2 048 features par image",
             size=Pt(14), color=RGBColor(0x88, 0x88, 0x88), new=True, space_before=Pt(12))

    # ═══════ SLIDE 7 — CONFIGURATION DU MODÈLE ═══════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 3, "Modèle ResNet50 — Configuration")

    add_table(
        s,
        headers=["Opération", "Code", "Effet"],
        rows=[
            ["Geler les poids", "requires_grad = False", "Aucun paramètre ne sera modifié"],
            ["Mode inférence", "model.eval()", "BatchNorm et Dropout désactivés"],
            ["Retirer la classification", "model.fc = nn.Identity()", "Sortie directe [batch, 2048]"],
            ["Pas de gradients", "torch.no_grad()", "Économie de mémoire"],
        ],
        top=Inches(1.8),
    )

    add_kpi_card(s, Inches(0.8), Inches(4.5), "23,5M", "Paramètres totaux", BLEU_ACCENT)
    add_kpi_card(s, Inches(3.6), Inches(4.5), "0", "Paramètres entraînables", VERT_OK)
    add_kpi_card(s, Inches(6.4), Inches(4.5), "2 048", "Features par image", OR_ACCENT)
    add_kpi_card(s, Inches(9.2), Inches(4.5), "CPU", "Device utilisé", GRIS_SUBTITLE)

    slide_footer(s)

    # ═══════ SLIDE 8 — RÉSULTATS DE L'EXTRACTION ═══════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 4, "Résultats de l'extraction")

    add_kpi_card(s, Inches(0.8), Inches(1.6), "1 506", "Images traitées", OR_ACCENT)
    add_kpi_card(s, Inches(3.6), Inches(1.6), "2 048", "Features par image", BLEU_ACCENT)
    add_kpi_card(s, Inches(6.4), Inches(1.6), "0 NaN", "Intégrité vérifiée", VERT_OK)
    add_kpi_card(s, Inches(9.2), Inches(1.6), "~108s", "Temps d'extraction (CPU)", GRIS_SUBTITLE)

    add_table(
        s,
        headers=["Métrique", "Valeur"],
        rows=[
            ["Shape", "(1 506, 2 048)"],
            ["Min", "0.0000"],
            ["Max", "8.1477"],
            ["Mean", "0.1007"],
            ["Std", "0.3043"],
            ["NaN", "0"],
            ["Inf", "0"],
        ],
        left=Inches(0.8), top=Inches(3.3), width=Inches(5),
    )

    add_content_block(s, "Sauvegarde", [
        "features.npy — Matrice numpy (1 506, 2 048)",
        "metadata.csv — DataFrame à 6 colonnes",
        "Stockés dans data/features/",
    ], left=Inches(7.0), top=Inches(3.3), width=Inches(5.5))

    slide_footer(s)

    # ═══════ SLIDE 9 — TABLEAU EXPLOITABLE ═══════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 5, "Tableau exploitable — 1 506 × 2 054")

    add_content_block(s, "Structure du tableau", [
        "6 colonnes de métadonnées : path, filename, dossier, label_name, label_id, is_labeled",
        "2 048 colonnes de features : feature_0 … feature_2047",
        "Total : 1 506 lignes × 2 054 colonnes",
    ], left=Inches(0.8), top=Inches(1.6), width=Inches(5.5))

    add_table(
        s,
        headers=["Colonne", "Type", "Exemple"],
        rows=[
            ["path", "str", "data/.../cancer/05340cd4.jpg"],
            ["filename", "str", "05340cd4-3bb2-…-8351.jpg"],
            ["dossier", "str", "cancer / normal / sans_label"],
            ["label_name", "str / None", "cancer / normal / None"],
            ["label_id", "Int64", "0 / 1 / NA"],
            ["is_labeled", "bool", "True / False"],
            ["feature_0…2047", "float64", "0.0000 … 8.1477"],
        ],
        left=Inches(7.0), top=Inches(1.6), width=Inches(5.8),
    )

    tf = add_textbox(s, Inches(0.8), Inches(4.5), Inches(11), Inches(1))
    set_para(tf, "→  Ce tableau est le résultat attendu de l'étape 2 : un « tableau exploitable » prêt pour le clustering (étape 3) et l'apprentissage semi-supervisé (étape 4).",
             size=Pt(14), color=GRIS_SUBTITLE)

    slide_footer(s)

    # ═══════ SLIDE 10 — SÉPARATEUR BILAN ═══════
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, BLEU_PROFOND)
    add_rect(s, 0, 0, SLIDE_W, Inches(0.06), OR_ACCENT)
    tf = add_textbox(s, Inches(1.2), Inches(2.5), Inches(10), Inches(1))
    set_para(tf, "BILAN", size=Pt(44), color=OR_ACCENT, bold=True)
    tf2 = add_textbox(s, Inches(1.2), Inches(3.8), Inches(10), Inches(1))
    set_para(tf2, "Résultats et prochaines étapes", size=Pt(24), color=BLANC)
    set_para(tf2, "Validation de l'extraction et préparation du clustering",
             size=Pt(14), color=RGBColor(0x88, 0x88, 0x88), new=True, space_before=Pt(12))

    # ═══════ SLIDE 11 — SYNTHÈSE & PROCHAINES ÉTAPES ═══════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 6, "Synthèse & Prochaines étapes")

    add_content_block(s, "Ce qui a été fait", [
        "Pipeline de preprocessing officiel ResNet50 appliqué aux 1 506 images",
        "Vérification visuelle avec OpenCV (brute vs preprocessed)",
        "Extraction de 2 048 features par image via avgpool de ResNet50",
        "Sauvegarde en features.npy + metadata.csv",
        "Tableau exploitable de 1 506 × 2 054 construit et affiché",
    ], left=Inches(0.8), top=Inches(1.6), width=Inches(5.5))

    add_content_block(s, "Prochaines étapes — Étape 3", [
        "Réduction de dimension (PCA, t-SNE) sur les 2 048 features",
        "Clustering exploratoire (K-Means, etc.) pour identifier des regroupements naturels",
        "Évaluation de la séparabilité cancer/normal dans l'espace des features",
    ], left=Inches(7.0), top=Inches(1.6), width=Inches(5.5))

    tf = add_textbox(s, Inches(0.8), Inches(4.5), Inches(11.5), Inches(2.5))
    set_para(tf, "Definition of done — Étape 2", size=Pt(16), color=BLEU_PRINCIPAL, bold=True,
             space_after=Pt(8))
    add_check_item(tf, "1 506 images chargées avec le même pipeline, sans erreur")
    add_check_item(tf, "Embeddings extraits via avgpool de ResNet50")
    add_check_item(tf, "Dimension vérifiée : (1 506, 2 048)")
    add_check_item(tf, "Aucun NaN ni Inf")
    add_check_item(tf, "Métadonnées à 6 colonnes reliées à chaque feature")
    add_check_item(tf, "Sorties sauvegardées (features.npy + metadata.csv)")
    add_check_item(tf, "Choix techniques justifiés dans le notebook")

    slide_footer(s)

    return prs


# ──────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────

if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(str(OUTPUT_FILE))
    print(f"✓ Présentation générée : {OUTPUT_FILE}")
    print(f"  Nombre de slides : {len(prs.slides)}")
