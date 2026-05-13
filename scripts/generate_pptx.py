"""
Script de génération du support de présentation — Projet 10
Labellisez et appliquez des approches semi-supervisées en traitement d'images

Auteur : David MEDRAGH
Date : Mai 2026

Description :
    Ce script génère un fichier PowerPoint (.pptx) unique couvrant toutes les
    étapes du projet. Il est conçu pour être enrichi au fur et à mesure
    de l'avancement (étapes 1 à 4). Limite : 15 slides max (consigne école).

    Le template visuel reprend le style du projet 9 (palette bleu profond / or,
    barre de titre numérotée, footer avec logos, KPI cards, séparateurs).

Slides produites (étapes 1-2) :
     1. Couverture
     2. Sommaire
     3. Séparateur ÉTAPE 1 — EXPLORATION
     4. Contexte & Dataset
     5. Vérification technique & Écarts
     6. Échantillon visuel
     7. Séparateur ÉTAPE 2 — FEATURES
     8. Pipeline de preprocessing
     9. Modèle ResNet50 — Configuration
    10. Résultats de l'extraction
    11. Séparateur BILAN
    12. Synthèse & Prochaines étapes

Usage :
    uv run python scripts/generate_pptx.py

Sortie :
    livrables/projet10_presentation.pptx
"""

import warnings
warnings.filterwarnings("ignore")

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- CHEMINS ---
PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT_ROOT / "livrables"
OUTPUT_FILE = OUTPUT_DIR / "projet10_presentation.pptx"
LOGO_CURELYTICSAI = PROJECT_ROOT / "doc" / "logo" / "logo_CurelyticsIA.png"
LOGO_OC = PROJECT_ROOT / "doc" / "logo" / "logo_openclassrooms.png"
ECHANTILLON_PATH = PROJECT_ROOT / "doc" / "dataset_preview" / "echantillon_visuel_etape1.png"
STACK_TECHNIQUE_PATH = PROJECT_ROOT / "doc" / "png" / "stack_technique_slide_etape1.png"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# --- PALETTE ---
BLEU_PROFOND = RGBColor(0x0D, 0x1B, 0x2A)
BLEU_PRINCIPAL = RGBColor(0x1B, 0x3A, 0x5C)
BLEU_ACCENT = RGBColor(0x3A, 0x7C, 0xBD)
OR_ACCENT = RGBColor(0xD4, 0xA0, 0x3C)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
GRIS_98 = RGBColor(0xFA, 0xFA, 0xFA)
GRIS_TEXTE = RGBColor(0x2D, 0x2D, 0x2D)
GRIS_SUBTITLE = RGBColor(0x6B, 0x6B, 0x6B)
GRIS_FOOTER = RGBColor(0xAA, 0xAA, 0xAA)
GRIS_LIGNE = RGBColor(0xDD, 0xDD, 0xDD)
VERT_OK = RGBColor(0x27, 0xAE, 0x60)
ORANGE_WARN = RGBColor(0xE6, 0x7E, 0x22)


# ──────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────

def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.fill.solid()
        shape.line.fill.fore_color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf


def set_para(tf, text, size=Pt(14), color=GRIS_TEXTE, bold=False,
             align=PP_ALIGN.LEFT, space_before=Pt(0), space_after=Pt(0), new=False):
    if new:
        p = tf.add_paragraph()
    else:
        p = (tf.paragraphs[0]
             if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == ""
             else tf.add_paragraph())
    run = p.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    p.alignment = align
    p.space_before = space_before
    p.space_after = space_after
    return p


def slide_title_bar(slide, number, title):
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.3), BLEU_PROFOND)
    add_rect(slide, 0, 0, Inches(0.08), Inches(1.3), OR_ACCENT)
    tf = add_textbox(slide, Inches(0.5), Inches(0.25), Inches(0.8), Inches(0.8))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf, f"{number:02d}", size=Pt(36), color=OR_ACCENT, bold=True)
    add_rect(slide, Inches(1.35), Inches(0.3), Pt(2), Inches(0.7), OR_ACCENT)
    tf = add_textbox(slide, Inches(1.6), Inches(0.25), Inches(10), Inches(0.8))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf, title, size=Pt(26), color=BLANC, bold=True)


def slide_footer(slide):
    add_rect(slide, 0, Inches(7.1), SLIDE_W, Pt(1), GRIS_LIGNE)
    tf = add_textbox(slide, Inches(0.5), Inches(7.15), Inches(6), Inches(0.3))
    set_para(tf, "CurelyticsIA — Projet BrainScanAI", size=Pt(9), color=GRIS_FOOTER)
    tf2 = add_textbox(slide, Inches(6.5), Inches(7.15), Inches(2), Inches(0.3))
    set_para(tf2, "Projet 10", size=Pt(9), color=BLEU_ACCENT, bold=True,
             align=PP_ALIGN.CENTER)
    tf3 = add_textbox(slide, Inches(9), Inches(7.15), Inches(3.8), Inches(0.3))
    set_para(tf3, "David Medragh — Data Scientist", size=Pt(9), color=GRIS_FOOTER,
             align=PP_ALIGN.RIGHT)
    logo_h = Inches(0.3)
    if LOGO_OC.exists():
        slide.shapes.add_picture(str(LOGO_OC), Inches(4.8), Inches(7.13), height=logo_h)
    if LOGO_CURELYTICSAI.exists():
        slide.shapes.add_picture(str(LOGO_CURELYTICSAI), Inches(5.5), Inches(7.13), height=logo_h)


def add_separator(prs, blank, title, subtitle, detail=None, image_path=None):
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, BLEU_PROFOND)
    add_rect(s, 0, 0, SLIDE_W, Inches(0.06), OR_ACCENT)
    tf = add_textbox(s, Inches(1.2), Inches(2.2), Inches(6), Inches(1))
    set_para(tf, title, size=Pt(44), color=OR_ACCENT, bold=True)
    tf2 = add_textbox(s, Inches(1.2), Inches(3.5), Inches(6), Inches(1))
    set_para(tf2, subtitle, size=Pt(24), color=BLANC)
    if detail:
        set_para(tf2, detail, size=Pt(14), color=RGBColor(0x88, 0x88, 0x88),
                 new=True, space_before=Pt(12))
    if image_path and Path(image_path).exists():
        s.shapes.add_picture(str(image_path), Inches(7.5), Inches(1.2), height=Inches(5.5))
    return s


def add_content_block(slide, title, bullets, left=Inches(0.8), top=Inches(1.6),
                      width=Inches(5.5)):
    tf = add_textbox(slide, left, top, width, Inches(5))
    set_para(tf, title, size=Pt(16), color=BLEU_PRINCIPAL, bold=True, space_after=Pt(8))
    for bullet in bullets:
        set_para(tf, f"▸  {bullet}", size=Pt(13), color=GRIS_TEXTE, space_before=Pt(5), new=True)
    return tf


def add_kpi_card(slide, left, top, value, label, value_color=BLEU_ACCENT,
                 w=Inches(2.5), h=Inches(1.3)):
    add_rect(slide, left, top, w, h, GRIS_98, GRIS_LIGNE)
    tf = add_textbox(slide, left + Inches(0.15), top + Inches(0.1),
                     w - Inches(0.3), Inches(0.7))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf, value, size=Pt(28), color=value_color, bold=True, align=PP_ALIGN.CENTER)
    tf2 = add_textbox(slide, left + Inches(0.15), top + Inches(0.75),
                      w - Inches(0.3), Inches(0.5))
    set_para(tf2, label, size=Pt(11), color=GRIS_SUBTITLE, align=PP_ALIGN.CENTER)


def add_check_item(tf, text, ok=True):
    prefix = "✓" if ok else "⚠"
    color = VERT_OK if ok else ORANGE_WARN
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = f"  {prefix}  {text}"
    run.font.size = Pt(14)
    run.font.color.rgb = color
    p.space_before = Pt(6)


def add_table(slide, headers, rows, left=Inches(0.8), top=Inches(1.8), width=Inches(11.5)):
    nb_rows = len(rows) + 1
    nb_cols = len(headers)
    table_shape = slide.shapes.add_table(
        nb_rows, nb_cols, left, top, width, Inches(0.38 * nb_rows))
    table = table_shape.table

    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLEU_PROFOND
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(12)
            paragraph.font.bold = True
            paragraph.font.color.rgb = BLANC
            paragraph.font.name = "Calibri"
            paragraph.alignment = PP_ALIGN.CENTER

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLANC if i % 2 == 0 else GRIS_98
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(11)
                paragraph.font.color.rgb = GRIS_TEXTE
                paragraph.font.name = "Calibri"
                paragraph.alignment = PP_ALIGN.CENTER


# ──────────────────────────────────────────────
# Construction des slides
# ──────────────────────────────────────────────

def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ═══════════════════════════════════════════
    # SLIDE 1 — COUVERTURE
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, BLEU_PROFOND)
    add_rect(s, 0, 0, SLIDE_W, Inches(0.06), OR_ACCENT)
    add_rect(s, 0, Inches(7.44), SLIDE_W, Inches(0.06), OR_ACCENT)

    tf = add_textbox(s, Inches(1.2), Inches(1.5), Inches(10), Inches(1.5))
    set_para(tf, "ANALYSE D'IRM CÉRÉBRALES", size=Pt(44), color=BLANC, bold=True)
    set_para(tf, "APPROCHES SEMI-SUPERVISÉES", size=Pt(44), color=OR_ACCENT, bold=True, new=True)

    add_rect(s, Inches(1.2), Inches(3.6), Inches(3), Pt(3), OR_ACCENT)

    tf2 = add_textbox(s, Inches(1.2), Inches(3.9), Inches(10), Inches(1))
    set_para(tf2, "Projet BrainScanAI — CurelyticsIA",
             size=Pt(20), color=RGBColor(0xBB, 0xBB, 0xBB))
    set_para(tf2, "Exploration, extraction de features, clustering et semi-supervisé",
             size=Pt(15), color=RGBColor(0x88, 0x88, 0x88), new=True, space_before=Pt(6))

    tf3 = add_textbox(s, Inches(1.2), Inches(5.3), Inches(10), Inches(1.5))
    set_para(tf3, "David MEDRAGH", size=Pt(18), color=BLANC, bold=True)
    set_para(tf3, "Data Scientist — CurelyticsIA", size=Pt(14),
             color=RGBColor(0x88, 0x88, 0x88), new=True, space_before=Pt(4))
    set_para(tf3, "Mai 2026", size=Pt(12),
             color=RGBColor(0x66, 0x66, 0x66), new=True, space_before=Pt(12))

    logo_h = Inches(0.6)
    if LOGO_CURELYTICSAI.exists():
        s.shapes.add_picture(str(LOGO_CURELYTICSAI), Inches(10.5), Inches(6.2), height=logo_h)
    if LOGO_OC.exists():
        s.shapes.add_picture(str(LOGO_OC), Inches(11.8), Inches(6.2), height=logo_h)

    # ═══════════════════════════════════════════
    # SLIDE 2 — SOMMAIRE
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SLIDE_W, Inches(1.3), BLEU_PROFOND)
    add_rect(s, 0, 0, Inches(0.08), Inches(1.3), OR_ACCENT)
    tf = add_textbox(s, Inches(0.5), Inches(0.25), Inches(10), Inches(0.8))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_para(tf, "Sommaire", size=Pt(26), color=BLANC, bold=True)

    x_positions = [Inches(0.5), Inches(3.7), Inches(6.9), Inches(10.1)]
    sommaire_cols = [
        ("Étape 1", "Exploration", [
            "Contexte & Dataset",
            "Vérification technique",
            "Échantillon visuel",
        ]),
        ("Étape 2", "Features", [
            "Pipeline preprocessing",
            "Modèle ResNet50",
            "Résultats extraction",
        ]),
        ("Étape 3", "Clustering", [
            "À venir",
        ]),
        ("Étape 4", "Semi-supervisé", [
            "À venir",
        ]),
    ]

    for col, (etape, subtitle, items) in enumerate(sommaire_cols):
        x = x_positions[col]
        y_start = Inches(1.6)
        tf = add_textbox(s, x, y_start, Inches(3), Inches(0.5))
        set_para(tf, etape, size=Pt(18), color=OR_ACCENT, bold=True)
        set_para(tf, subtitle, size=Pt(12), color=GRIS_SUBTITLE, new=True,
                 space_before=Pt(4))
        tf2 = add_textbox(s, x + Inches(0.2), y_start + Inches(0.9), Inches(2.8),
                          Inches(4.5))
        for i, item in enumerate(items):
            color = GRIS_TEXTE if item != "À venir" else GRIS_SUBTITLE
            set_para(tf2, f"▸  {item}", size=Pt(11), color=color,
                     space_before=Pt(6), new=(i > 0))

    slide_footer(s)

    # ═══════════════════════════════════════════
    # SLIDE 3 — SÉPARATEUR ÉTAPE 1
    # ═══════════════════════════════════════════
    add_separator(prs, blank,
                  "ÉTAPE 1 — EXPLORATION",
                  "Contexte, dataset et vérification technique",
                  "Inventaire complet, contrôle qualité et écarts documentés",
                  STACK_TECHNIQUE_PATH)

    # ═══════════════════════════════════════════
    # SLIDE 4 — CONTEXTE & DATASET
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 1, "Contexte & Dataset")

    add_content_block(s, "CurelyticsIA", [
        "Startup e-santé — analyse d'images médicales par IA",
        "Projet R&D BrainScanAI : détection de tumeurs cérébrales",
        "Dataset d'IRM collecté auprès de radiologues experts",
    ], left=Inches(0.8), top=Inches(1.6), width=Inches(5.5))

    add_kpi_card(s, Inches(7.0), Inches(1.6), "1 506", "Images totales", OR_ACCENT)
    add_kpi_card(s, Inches(10.0), Inches(1.6), "100", "Labellisées (6,6 %)", VERT_OK)

    add_content_block(s, "Structure", [
        "avec_labels/cancer/    →  50 IRM avec tumeurs",
        "avec_labels/normal/    →  50 IRM cerveaux sains",
        "sans_label/                 → 1 406 IRM non étiquetées",
    ], left=Inches(7.0), top=Inches(3.3), width=Inches(5.5))

    add_kpi_card(s, Inches(0.8), Inches(4.5), "512×512", "Résolution (px)", BLEU_PRINCIPAL,
                 w=Inches(2.2))
    add_kpi_card(s, Inches(3.3), Inches(4.5), "JPEG", "Format", BLEU_ACCENT,
                 w=Inches(2.2))
    add_kpi_card(s, Inches(0.8), Inches(6.0), "RGB", "Mode couleur", BLEU_ACCENT,
                 w=Inches(2.2))
    add_kpi_card(s, Inches(3.3), Inches(6.0), "~6,6 %", "Labels disponibles", ORANGE_WARN,
                 w=Inches(2.2))

    slide_footer(s)

    # ═══════════════════════════════════════════
    # SLIDE 5 — VÉRIFICATION TECHNIQUE & ÉCARTS
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 2, "Vérification technique & Écarts")

    tf = add_textbox(s, Inches(0.8), Inches(1.5), Inches(5.5), Inches(4.5))
    set_para(tf, "Résultats du scan complet (1 506 images)", size=Pt(16),
             color=BLEU_PRINCIPAL, bold=True, space_after=Pt(12))
    add_check_item(tf, "Toutes les images lisibles — aucune corrompue")
    add_check_item(tf, "Résolution homogène : 512×512")
    add_check_item(tf, "Format uniformément JPEG")
    add_check_item(tf, "Mode couleur uniformément RGB")
    add_check_item(tf, "Aucun doublon entre les 3 dossiers")

    add_table(
        s,
        headers=["Critère", "Énoncé", "Réel", "Commentaire"],
        rows=[
            ["Format", "PNG (courriel)", "JPEG (.jpg)", "Fichier descriptif confirme JPEG"],
            ["Non étiquetées", "1 400", "1 406", "+6 images"],
            ["Total", "1 500", "1 506", "Écart cohérent"],
            ["Métadonnées", "Non mentionnées", "Absentes", "Dossiers = labels"],
        ],
        left=Inches(6.8), top=Inches(1.8), width=Inches(6),
    )

    tf2 = add_textbox(s, Inches(0.8), Inches(5.5), Inches(11), Inches(1))
    set_para(tf2, "→  Écarts mineurs et documentés. Dataset propre, exploitable en l'état.",
             size=Pt(14), color=GRIS_SUBTITLE)

    slide_footer(s)

    # ═══════════════════════════════════════════
    # SLIDE 6 — ÉCHANTILLON VISUEL
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 3, "Échantillon visuel — 5 images par catégorie")

    if ECHANTILLON_PATH.exists():
        s.shapes.add_picture(
            str(ECHANTILLON_PATH),
            Inches(1.5), Inches(1.5),
            width=Inches(10.3),
        )
    else:
        tf = add_textbox(s, Inches(2), Inches(3), Inches(9), Inches(1))
        set_para(tf, "⚠  Exécuter le notebook étape 1 pour générer l'échantillon visuel.",
                 size=Pt(16), color=ORANGE_WARN)

    slide_footer(s)

    # ═══════════════════════════════════════════
    # SLIDE 7 — SÉPARATEUR ÉTAPE 2
    # ═══════════════════════════════════════════
    add_separator(prs, blank,
                  "ÉTAPE 2 — FEATURES",
                  "Des images brutes aux embeddings visuels",
                  "Preprocessing officiel ResNet50 — 2 048 features par image")

    # ═══════════════════════════════════════════
    # SLIDE 8 — PIPELINE DE PREPROCESSING
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 4, "Pipeline de preprocessing")

    add_content_block(s, "Preprocessing officiel ResNet50", [
        "weights = ResNet50_Weights.DEFAULT",
        "preprocess = weights.transforms()",
        "Pipeline unique — pas de transforms parallèle",
    ], left=Inches(0.8), top=Inches(1.6), width=Inches(5.5))

    add_content_block(s, "Étapes du pipeline", [
        "Resize(232) — côté court",
        "CenterCrop(224) — recadrage central 224×224",
        "ToTensor() → tenseur [0, 1]",
        "Normalize(mean ImageNet, std ImageNet)",
    ], left=Inches(7.0), top=Inches(1.6), width=Inches(5.5))

    add_content_block(s, "Dataset unifié", [
        "1 506 images dans un DataFrame unique (6 colonnes de métadonnées)",
        "DataLoader shuffle=False — ordre stable garanti",
        "Vérification visuelle avec OpenCV (brute vs preprocessed)",
    ], left=Inches(0.8), top=Inches(4.2), width=Inches(5.5))

    add_content_block(s, "Outils de l'école", [
        "torchvision — preprocessing officiel + modèle",
        "OpenCV (cv2) — chargement brut + conversion BGR→RGB",
        "PIL (Pillow) — ouverture dans le Dataset custom",
        "numpy, pandas, matplotlib",
    ], left=Inches(7.0), top=Inches(4.2), width=Inches(5.5))

    slide_footer(s)

    # ═══════════════════════════════════════════
    # SLIDE 9 — MODÈLE RESNET50
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 5, "Modèle ResNet50 — Configuration")

    add_table(
        s,
        headers=["Opération", "Code", "Effet"],
        rows=[
            ["Geler les poids", "requires_grad = False", "Paramètres non modifiés"],
            ["Mode inférence", "model.eval()", "BatchNorm / Dropout désactivés"],
            ["Retirer la classification", "model.fc = nn.Identity()", "Sortie [batch, 2048]"],
            ["Pas de gradients", "torch.no_grad()", "Économie de mémoire"],
        ],
        top=Inches(1.8),
    )

    add_kpi_card(s, Inches(0.8), Inches(4.5), "23,5M", "Paramètres totaux", BLEU_ACCENT)
    add_kpi_card(s, Inches(3.6), Inches(4.5), "0", "Entraînables", VERT_OK)
    add_kpi_card(s, Inches(6.4), Inches(4.5), "2 048", "Features / image", OR_ACCENT)
    add_kpi_card(s, Inches(9.2), Inches(4.5), "CPU", "Device", GRIS_SUBTITLE)

    slide_footer(s)

    # ═══════════════════════════════════════════
    # SLIDE 10 — RÉSULTATS DE L'EXTRACTION
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 6, "Résultats de l'extraction")

    add_kpi_card(s, Inches(0.8), Inches(1.6), "1 506", "Images traitées", OR_ACCENT)
    add_kpi_card(s, Inches(3.6), Inches(1.6), "2 048", "Features / image", BLEU_ACCENT)
    add_kpi_card(s, Inches(6.4), Inches(1.6), "0 NaN", "Intégrité OK", VERT_OK)
    add_kpi_card(s, Inches(9.2), Inches(1.6), "~108s", "Temps (CPU)", GRIS_SUBTITLE)

    add_table(
        s,
        headers=["Métrique", "Valeur"],
        rows=[
            ["Shape", "(1 506, 2 048)"],
            ["Min", "0.0000"],
            ["Max", "8.1477"],
            ["Mean", "0.1007"],
            ["Std", "0.3043"],
        ],
        left=Inches(0.8), top=Inches(3.3), width=Inches(5),
    )

    add_content_block(s, "Sauvegarde", [
        "features.npy — matrice (1 506, 2 048)",
        "metadata.csv — 6 colonnes de métadonnées",
        "Tableau exploitable : 1 506 × 2 054 colonnes",
    ], left=Inches(7.0), top=Inches(3.3), width=Inches(5.5))

    add_content_block(s, "Métadonnées", [
        "path, filename, dossier",
        "label_name, label_id (Int64), is_labeled",
    ], left=Inches(7.0), top=Inches(5.3), width=Inches(5.5))

    slide_footer(s)

    # ═══════════════════════════════════════════
    # SLIDE 11 — SÉPARATEUR BILAN
    # ═══════════════════════════════════════════
    add_separator(prs, blank,
                  "BILAN",
                  "Résultats et prochaines étapes",
                  "Validation des étapes 1-2, préparation du clustering")

    # ═══════════════════════════════════════════
    # SLIDE 12 — SYNTHÈSE & PROCHAINES ÉTAPES
    # ═══════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    slide_title_bar(s, 7, "Synthèse & Prochaines étapes")

    add_content_block(s, "Étape 1 — Exploration ✓", [
        "Dataset propre, homogène, exploitable en l'état",
        "1 506 images vérifiées — aucune anomalie bloquante",
        "Écarts documentés (format, nombre d'images)",
        "Fort déséquilibre : ~6,6 % de labels",
    ], left=Inches(0.8), top=Inches(1.6), width=Inches(5.5))

    add_content_block(s, "Étape 2 — Features ✓", [
        "Preprocessing officiel ResNet50 appliqué",
        "2 048 features extraites pour chaque image",
        "0 NaN, 0 Inf — extraction propre",
        "Sauvegarde features.npy + metadata.csv",
    ], left=Inches(7.0), top=Inches(1.6), width=Inches(5.5))

    add_content_block(s, "Prochaines étapes", [
        "Étape 3 — Réduction de dimension + clustering exploratoire",
        "Étape 4 — Apprentissage semi-supervisé avec labels partiels",
        "Recommandations pour passage à l'échelle (5 000 € / 4M images)",
    ], left=Inches(0.8), top=Inches(4.5), width=Inches(11))

    slide_footer(s)

    return prs


# ──────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────

if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(str(OUTPUT_FILE))
    print(f"✓ Présentation générée : {OUTPUT_FILE}")
    print(f"  Nombre de slides : {len(prs.slides)}")
